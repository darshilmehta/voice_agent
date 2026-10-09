# ruff: noqa: E501  (document text and evidence strings are data)
"""Document 2: Zephyra Logistics Limited, Q4 and FY24 investor presentation (PPTX, 15 slides).

A second company with overlapping metric names (revenue, EBITDA margin, PAT, net debt, employees, a Digital Services
segment) next to Valmora's annual report, so the two documents have to be told apart. Slide titles name the company
on some slides and not on others, as real decks do.
"""

from __future__ import annotations

from pathlib import Path

from app.evals.manifest import DocumentEntry

from .finance import Zephyra as Z
from .finance import bps, growth, n0
from .registry import F, FactSpec, Registry

NAME = "zephyra_investor_deck_q4fy24.pptx"
TITLE = "Zephyra Logistics Limited - Investor presentation, Q4 and FY24"
SLIDES = 15


class Slide:
    def __init__(self, deck: Deck, number: int, layout: int, title: str) -> None:
        self.deck = deck
        self.number = number
        self.slide = deck.prs.slides.add_slide(deck.prs.slide_layouts[layout])
        self.slide.shapes.title.text = title
        self._top = 1.55  # inches, where the next block starts

    def _plant(self, text: str, facts: tuple[FactSpec, ...]) -> None:
        for spec in facts:
            self.deck.reg.plant(spec, document=NAME, text=text, page=self.number)

    def bullets(
        self, items: list[str | tuple[str, tuple[FactSpec, ...]]], *, size: int = 18, height: float = 4.8
    ) -> None:
        from pptx.util import Inches, Pt

        box = self.slide.shapes.add_textbox(Inches(0.6), Inches(self._top), Inches(8.8), Inches(height))
        tf = box.text_frame
        tf.word_wrap = True
        first = True
        for item in items:
            text, facts = (item, ()) if isinstance(item, str) else item
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.text = "• " + text
            p.space_after = Pt(8)
            for run in p.runs:
                run.font.size = Pt(size)
            self._plant(text, facts)
        self._top += height

    def text(self, text: str, *facts: FactSpec, size: int = 16, height: float = 0.8) -> None:
        self.bullets([(text, facts)], size=size, height=height)

    def table(
        self,
        rows: list[list[str]],
        widths: list[float],
        *,
        facts: dict[int, tuple[FactSpec, ...] | FactSpec] | None = None,
        size: int = 13,
        row_h: float = 0.42,
    ) -> None:
        from pptx.util import Inches, Pt

        total_w = 8.8
        shape = self.slide.shapes.add_table(
            len(rows), len(rows[0]), Inches(0.6), Inches(self._top), Inches(total_w), Inches(row_h * len(rows))
        )
        table = shape.table
        for c, w in enumerate(widths):
            table.columns[c].width = Inches(total_w * w)
        for r, row in enumerate(rows):
            for c, cell in enumerate(row):
                table.cell(r, c).text = cell
                for p in table.cell(r, c).text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(size)
        for r, specs in (facts or {}).items():
            specs = (specs,) if isinstance(specs, FactSpec) else specs
            row_text = " ".join(rows[r])
            for spec in specs:
                self.deck.reg.plant(spec, document=NAME, text=row_text, page=self.number)
        self._top += row_h * len(rows) + 0.3


class Deck:
    def __init__(self, reg: Registry) -> None:
        from pptx import Presentation

        self.reg = reg
        self.prs = Presentation()  # default 4:3 template: the layouts' placeholders fit without repositioning
        self.slides: list[Slide] = []

    def slide(self, title: str, layout: int = 5) -> Slide:
        s = Slide(self, len(self.slides) + 1, layout, title)
        self.slides.append(s)
        return s


def build(reg: Registry, out_dir: Path) -> DocumentEntry:
    entry = DocumentEntry(name=NAME, title=TITLE, format="pptx", language="en", pages=SLIDES, company=Z.name)
    reg.add_document(entry)
    deck = Deck(reg)
    for make in (s01, s02, s03, s04, s05, s06, s07, s08, s09, s10, s11, s12, s13, s14, s15):
        make(deck)
    assert len(deck.slides) == SLIDES
    deck.prs.core_properties.title = TITLE
    deck.prs.core_properties.author = "Eval corpus generator (fictional company)"
    deck.prs.save(str(out_dir / NAME))
    return entry


def cr(x: float) -> str:
    return f"₹ {n0(x)} crore"


# ------------------------------------------------------------------ slides


def s01(d: Deck) -> None:
    s = d.slide("Zephyra Logistics Limited", layout=0)
    sub = s.slide.placeholders[1]
    sub.text = "Investor presentation: Q4 and FY24 results\nMay 2024"
    d.reg.plant(
        F(
            "zph.deck_date",
            ["Q4 and FY24 results", "May 2024"],
            "The deck presents Q4 and FY24 results, dated May 2024",
            "text",
        ),
        document=NAME,
        text=sub.text,
        page=1,
    )


def s02(d: Deck) -> None:
    s = d.slide("Safe harbour and agenda")
    s.bullets(
        [
            "Statements about the future in this presentation are based on current expectations and are subject to risks and uncertainties.",
            "Agenda: FY24 at a glance; segment performance; financial results; network and ESG; outlook; stock information.",
            (
                "The Q4 FY24 earnings call is on 16 May 2024 at 4:00 p.m. IST.",
                (
                    F(
                        "zph.call",
                        ["earnings call", "16 May 2024", "4:00 p.m. IST"],
                        "Zephyra's Q4 FY24 earnings call is on 16 May 2024 at 4:00 p.m. IST",
                        "date",
                    ),
                ),
            ),
        ],
        size=17,
    )


def s03(d: Deck) -> None:
    s = d.slide("Zephyra at a glance: FY24")
    r24, r23, e24, e23 = Z.revenue["FY24"], Z.revenue["FY23"], Z.ebitda["FY24"], Z.ebitda["FY23"]
    rows = [
        ["Metric", "FY24", "FY23", "Change"],
        ["Revenue from operations (₹ crore)", n0(r24), n0(r23), "+" + growth(r24, r23)],
        ["EBITDA (₹ crore)", n0(e24), n0(e23), "+" + growth(e24, e23)],
        ["EBITDA margin", Z.margin("FY24"), Z.margin("FY23"), f"+{bps(e24, r24, e23, r23)} bps"],
        [
            "Profit after tax (₹ crore)",
            n0(Z.pat["FY24"]),
            n0(Z.pat["FY23"]),
            "+" + growth(Z.pat["FY24"], Z.pat["FY23"]),
        ],
        [
            "Net debt (₹ crore)",
            n0(Z.net_debt["FY24"]),
            n0(Z.net_debt["FY23"]),
            growth(Z.net_debt["FY24"], Z.net_debt["FY23"]),
        ],
        ["Net debt to EBITDA", Z.nd_ebitda("FY24"), Z.nd_ebitda("FY23"), "-0.6x"],
        ["Return on capital employed", Z.roce["FY24"], Z.roce["FY23"], "+1.9 pts"],
        [
            "Number of employees",
            n0(Z.employees["FY24"]),
            n0(Z.employees["FY23"]),
            "+" + growth(Z.employees["FY24"], Z.employees["FY23"]),
        ],
    ]
    s.table(
        rows,
        [0.42, 0.18, 0.18, 0.22],
        facts={
            1: (
                F(
                    "zph.rev.fy24",
                    ["Revenue from operations", n0(r24)],
                    f"Zephyra's FY24 revenue from operations was ₹ {n0(r24)} crore",
                ),
                F(
                    "zph.rev.fy23",
                    ["Revenue from operations", n0(r23)],
                    f"Zephyra's FY23 revenue from operations was ₹ {n0(r23)} crore",
                ),
            ),
            2: F("zph.ebitda.fy24", ["EBITDA (", n0(e24)], f"Zephyra's FY24 EBITDA was ₹ {n0(e24)} crore"),
            3: (
                F(
                    "zph.margin.fy24",
                    ["EBITDA margin", Z.margin("FY24")],
                    f"Zephyra's FY24 EBITDA margin was {Z.margin('FY24')}",
                ),
                F(
                    "zph.margin.fy23",
                    ["EBITDA margin", Z.margin("FY23")],
                    f"Zephyra's FY23 EBITDA margin was {Z.margin('FY23')}",
                ),
            ),
            4: (
                F(
                    "zph.pat.fy24",
                    ["Profit after tax", n0(Z.pat["FY24"])],
                    f"Zephyra's FY24 profit after tax was ₹ {n0(Z.pat['FY24'])} crore",
                ),
                F(
                    "zph.pat.fy23",
                    ["Profit after tax", n0(Z.pat["FY23"])],
                    f"Zephyra's FY23 profit after tax was ₹ {n0(Z.pat['FY23'])} crore",
                ),
            ),
            5: F(
                "zph.net_debt",
                ["Net debt (", n0(Z.net_debt["FY24"]), n0(Z.net_debt["FY23"])],
                "Zephyra's net debt was ₹ 912 crore in FY24 (FY23: ₹ 1,075 crore)",
            ),
            6: F(
                "zph.nd_ebitda",
                ["Net debt to EBITDA", Z.nd_ebitda("FY24"), Z.nd_ebitda("FY23")],
                "Zephyra's net debt to EBITDA was 1.4x in FY24 (FY23: 2.0x)",
            ),
            7: F(
                "zph.roce",
                ["Return on capital employed", Z.roce["FY24"], Z.roce["FY23"]],
                "Zephyra's ROCE was 14.8% in FY24 (FY23: 12.9%)",
            ),
            8: F(
                "zph.employees",
                ["Number of employees", n0(Z.employees["FY24"])],
                f"Zephyra had {n0(Z.employees['FY24'])} employees in FY24 (FY23: {n0(Z.employees['FY23'])})",
            ),
        },
    )


def s04(d: Deck) -> None:
    s = d.slide("Our business model")
    s.bullets(
        [
            "Zephyra moves goods for more than 3,400 customers through three integrated businesses: Freight Services, Contract Logistics and Digital Services.",
            (
                "Founded in 2005 and listed in 2021; headquartered in Mumbai.",
                (
                    F(
                        "zph.founded",
                        ["Founded in 2005", "listed in 2021"],
                        "Zephyra was founded in 2005 and listed in 2021",
                        "number",
                    ),
                ),
            ),
            (
                "Asset-light where possible: 38 per cent of freight capacity is contracted from partner carriers.",
                (
                    F(
                        "zph.asset_light",
                        ["38 per cent of freight capacity"],
                        "38% of Zephyra's freight capacity is contracted from partners",
                        "number",
                    ),
                ),
            ),
            "A single control-tower platform gives customers visibility of every shipment from pick-up to proof of delivery.",
        ]
    )


def s05(d: Deck) -> None:
    s = d.slide("Segment overview")
    rows = [["Segment", "Revenue FY24", "Revenue FY23", "Growth", "EBITDA margin FY24", "EBITDA margin FY23"]]
    facts = {}
    keys = {"Freight Services": "freight", "Contract Logistics": "contract", "Digital Services": "digital"}
    for i, seg in enumerate(Z.segments, start=1):
        r24, r23 = Z.segments[seg]["FY24"][0], Z.segments[seg]["FY23"][0]
        rows.append(
            [seg, n0(r24), n0(r23), "+" + growth(r24, r23), Z.seg_margin(seg, "FY24"), Z.seg_margin(seg, "FY23")]
        )
        k = keys[seg]
        facts[i] = (
            F(
                f"zph.seg.{k}.rev.fy24",
                [seg, n0(r24)],
                f"Zephyra {seg} revenue was ₹ {n0(r24)} crore in FY24",
                "table_cell",
            ),
            F(
                f"zph.seg.{k}.margin.fy24",
                [seg, Z.seg_margin(seg, "FY24")],
                f"Zephyra {seg} EBITDA margin was {Z.seg_margin(seg, 'FY24')} in FY24",
                "table_cell",
            ),
        )
    rows.append(
        [
            "Total",
            n0(Z.revenue["FY24"]),
            n0(Z.revenue["FY23"]),
            "+" + growth(Z.revenue["FY24"], Z.revenue["FY23"]),
            Z.margin("FY24"),
            Z.margin("FY23"),
        ]
    )
    s.table(rows, [0.22, 0.16, 0.16, 0.12, 0.17, 0.17], facts=facts, size=12)
    s.text("Revenue in ₹ crore. Digital Services is the smallest and fastest-growing segment.", size=14)


def s06(d: Deck) -> None:
    s = d.slide("Freight Services")
    seg = "Freight Services"
    r24, r23 = Z.segments[seg]["FY24"][0], Z.segments[seg]["FY23"][0]
    s.bullets(
        [
            (
                f"Revenue of {cr(r24)}, up {growth(r24, r23)} on FY23.",
                (
                    F(
                        "zph.freight.rev",
                        [f"{n0(r24)} crore", growth(r24, r23)],
                        f"Freight Services revenue was ₹ {n0(r24)} crore, up {growth(r24, r23)}",
                        "number",
                    ),
                ),
            ),
            (f"EBITDA margin of {Z.seg_margin(seg, 'FY24')} (FY23: {Z.seg_margin(seg, 'FY23')}).", ()),
            (
                "Air freight volumes grew 14 per cent; the trucking fleet grew to 2,840 owned vehicles.",
                (
                    F(
                        "zph.freight.fleet",
                        ["2,840 owned vehicles", "14 per cent"],
                        "Zephyra's trucking fleet grew to 2,840 owned vehicles; air freight volumes grew 14%",
                        "number",
                    ),
                ),
            ),
            (
                "Average fleet age improved to 3.9 years.",
                (F("zph.freight.fleet_age", ["fleet age", "3.9 years"], "Average fleet age was 3.9 years", "number"),),
            ),
        ]
    )


def s07(d: Deck) -> None:
    s = d.slide("Contract Logistics")
    seg = "Contract Logistics"
    r24, r23 = Z.segments[seg]["FY24"][0], Z.segments[seg]["FY23"][0]
    s.bullets(
        [
            (
                f"Revenue of {cr(r24)}, up {growth(r24, r23)} on FY23.",
                (
                    F(
                        "zph.contract.rev",
                        [f"{n0(r24)} crore", growth(r24, r23)],
                        f"Contract Logistics revenue was ₹ {n0(r24)} crore, up {growth(r24, r23)}",
                        "number",
                    ),
                ),
            ),
            (f"EBITDA margin of {Z.seg_margin(seg, 'FY24')} (FY23: {Z.seg_margin(seg, 'FY23')}).", ()),
            (
                "Order book of ₹ 3,450 crore at 31 March 2024, signed for an average of 4.2 years.",
                (
                    F(
                        "zph.contract.orderbook",
                        ["Order book of ₹ 3,450 crore"],
                        "Contract Logistics order book was ₹ 3,450 crore at 31 March 2024",
                        "number",
                    ),
                ),
            ),
            (
                "Warehouse capacity of 18.6 million sq ft across 112 warehouses.",
                (
                    F(
                        "zph.contract.warehouses",
                        ["18.6 million sq ft", "112 warehouses"],
                        "Zephyra has 18.6 million sq ft of warehouse capacity across 112 warehouses",
                        "number",
                    ),
                ),
            ),
        ]
    )


def s08(d: Deck) -> None:
    s = d.slide("Digital Services")
    seg = "Digital Services"
    r24, r23 = Z.segments[seg]["FY24"][0], Z.segments[seg]["FY23"][0]
    s.bullets(
        [
            (
                f"Revenue of {cr(r24)}, up {growth(r24, r23)} on FY23.",
                (
                    F(
                        "zph.digital.rev",
                        [f"{n0(r24)} crore", growth(r24, r23)],
                        f"Zephyra Digital Services revenue was ₹ {n0(r24)} crore, up {growth(r24, r23)}",
                        "number",
                    ),
                ),
            ),
            (
                f"EBITDA margin of {Z.seg_margin(seg, 'FY24')} (FY23: {Z.seg_margin(seg, 'FY23')}).",
                (
                    F(
                        "zph.digital.margin",
                        [f"EBITDA margin of {Z.seg_margin(seg, 'FY24')}"],
                        f"Zephyra Digital Services EBITDA margin was {Z.seg_margin(seg, 'FY24')} in FY24",
                        "number",
                    ),
                ),
            ),
            (
                "The ZenTrack control-tower platform tracks 1.9 million shipments a month.",
                (
                    F(
                        "zph.digital.zentrack",
                        ["ZenTrack", "1.9 million shipments a month"],
                        "ZenTrack tracks 1.9 million shipments a month",
                        "number",
                    ),
                ),
            ),
            "Digital Services sells software to logistics customers and to Zephyra's own operating businesses.",
        ]
    )


def s09(d: Deck) -> None:
    s = d.slide("Q4 FY24 results")
    q = Z.q4
    rows = [
        ["Metric", "Q4 FY24", "Q4 FY23", "Q3 FY24"],
        ["Revenue from operations (₹ crore)", *[n0(q[k][0]) for k in ("Q4 FY24", "Q4 FY23", "Q3 FY24")]],
        ["EBITDA (₹ crore)", *[n0(q[k][1]) for k in ("Q4 FY24", "Q4 FY23", "Q3 FY24")]],
        ["EBITDA margin", *[Z.q_margin(k) for k in ("Q4 FY24", "Q4 FY23", "Q3 FY24")]],
        ["Profit after tax (₹ crore)", *[n0(q[k][2]) for k in ("Q4 FY24", "Q4 FY23", "Q3 FY24")]],
    ]
    s.table(
        rows,
        [0.43, 0.19, 0.19, 0.19],
        facts={
            1: F(
                "zph.q4.rev",
                ["Revenue from operations", n0(q["Q4 FY24"][0]), n0(q["Q4 FY23"][0])],
                "Zephyra's Q4 FY24 revenue was ₹ 1,329 crore (Q4 FY23: ₹ 1,186 crore)",
                "table_cell",
            ),
            3: F(
                "zph.q4.margin",
                ["EBITDA margin", Z.q_margin("Q4 FY24"), Z.q_margin("Q4 FY23")],
                f"Zephyra's Q4 FY24 EBITDA margin was {Z.q_margin('Q4 FY24')} (Q4 FY23: {Z.q_margin('Q4 FY23')})",
                "table_cell",
            ),
            4: F(
                "zph.q4.pat",
                ["Profit after tax", n0(q["Q4 FY24"][2]), n0(q["Q4 FY23"][2])],
                "Zephyra's Q4 FY24 profit after tax was ₹ 78 crore (Q4 FY23: ₹ 61 crore)",
                "table_cell",
            ),
        },
    )
    s.text(
        "The fourth quarter benefited from peak-season freight volumes and the start of two new warehouse contracts.",
        size=15,
        height=0.9,
    )


def s10(d: Deck) -> None:
    s = d.slide("FY24 profit and loss summary")
    # EBITDA + other income - depreciation - finance costs = profit before tax; tax at about 25 per cent
    da, fin, oi, tax = (
        {"FY24": 226, "FY23": 190},
        {"FY24": 90, "FY23": 88},
        {"FY24": 41, "FY23": 33},
        {"FY24": 90, "FY23": 71},
    )
    for y in ("FY24", "FY23"):
        assert Z.ebitda[y] + oi[y] - da[y] - fin[y] - tax[y] == Z.pat[y], y
    pbt = {y: Z.ebitda[y] + oi[y] - da[y] - fin[y] for y in ("FY24", "FY23")}
    rows = [
        ["₹ crore", "FY24", "FY23"],
        ["Revenue from operations", n0(Z.revenue["FY24"]), n0(Z.revenue["FY23"])],
        ["EBITDA", n0(Z.ebitda["FY24"]), n0(Z.ebitda["FY23"])],
        ["Other income", n0(oi["FY24"]), n0(oi["FY23"])],
        ["Depreciation and amortisation", n0(da["FY24"]), n0(da["FY23"])],
        ["Finance costs", n0(fin["FY24"]), n0(fin["FY23"])],
        ["Profit before tax", n0(pbt["FY24"]), n0(pbt["FY23"])],
        ["Tax expense", n0(tax["FY24"]), n0(tax["FY23"])],
        ["Profit after tax", n0(Z.pat["FY24"]), n0(Z.pat["FY23"])],
    ]
    s.table(
        rows,
        [0.5, 0.25, 0.25],
        facts={
            4: F(
                "zph.pl.da",
                ["Depreciation and amortisation", n0(da["FY24"]), n0(da["FY23"])],
                "Zephyra's depreciation and amortisation was ₹ 226 crore (FY24) and ₹ 190 crore (FY23)",
                "table_cell",
            ),
            5: F(
                "zph.pl.finance",
                ["Finance costs", n0(fin["FY24"]), n0(fin["FY23"])],
                "Zephyra's finance costs were ₹ 90 crore (FY24) and ₹ 88 crore (FY23)",
                "table_cell",
            ),
            6: F(
                "zph.pl.pbt",
                ["Profit before tax", n0(pbt["FY24"]), n0(pbt["FY23"])],
                "Zephyra's profit before tax was ₹ 358 crore (FY24) and ₹ 283 crore (FY23)",
                "table_cell",
            ),
        },
        row_h=0.4,
    )
    s.text(
        f"PAT margin was {Z.pat_margin('FY24')} in FY24 against {Z.pat_margin('FY23')} in FY23.",
        F(
            "zph.pl.pat_margin",
            ["PAT margin", Z.pat_margin("FY24"), Z.pat_margin("FY23")],
            "Zephyra's PAT margin was 5.4% in FY24 (FY23: 4.8%)",
            "number",
        ),
        size=15,
    )


def s11(d: Deck) -> None:
    s = d.slide("Balance sheet and cash flow")
    s.bullets(
        [
            (
                f"Net debt fell to {cr(Z.net_debt['FY24'])} from {cr(Z.net_debt['FY23'])}; net debt to EBITDA improved to {Z.nd_ebitda('FY24')} from {Z.nd_ebitda('FY23')}.",
                (
                    F(
                        "zph.bs.net_debt",
                        ["Net debt fell to", "912 crore", "1,075 crore"],
                        "Zephyra's net debt fell to ₹ 912 crore from ₹ 1,075 crore",
                        "number",
                    ),
                ),
            ),
            (
                f"Operating cash flow was {cr(Z.cfo['FY24'])} (FY23: {cr(Z.cfo['FY23'])}).",
                (
                    F(
                        "zph.bs.cfo",
                        ["Operating cash flow was", "604 crore", "497 crore"],
                        "Zephyra's operating cash flow was ₹ 604 crore in FY24 (FY23: ₹ 497 crore)",
                        "number",
                    ),
                ),
            ),
            (
                f"Capital expenditure was {cr(Z.capex['FY24'])} (FY23: {cr(Z.capex['FY23'])}), mainly warehouses and electric vehicles.",
                (
                    F(
                        "zph.bs.capex",
                        ["Capital expenditure was", "412 crore", "338 crore"],
                        "Zephyra's capex was ₹ 412 crore in FY24 (FY23: ₹ 338 crore)",
                        "number",
                    ),
                ),
            ),
            (
                f"Free cash flow rose to {cr(Z.fcf('FY24'))} from {cr(Z.fcf('FY23'))}.",
                (
                    F(
                        "zph.bs.fcf",
                        ["Free cash flow rose to", "192 crore", "159 crore"],
                        "Zephyra's free cash flow rose to ₹ 192 crore from ₹ 159 crore",
                        "number",
                    ),
                ),
            ),
            "Cash and bank balances were ₹ 287 crore at 31 March 2024.",
        ],
        size=17,
    )


def s12(d: Deck) -> None:
    s = d.slide("Network and fleet")
    s.table(
        [
            ["Network indicator", "FY24", "FY23"],
            ["Branches", "214", "198"],
            ["Warehouses", "112", "103"],
            ["Warehouse capacity (million sq ft)", "18.6", "16.9"],
            ["Owned trucks", "2,840", "2,510"],
            ["On-time delivery rate", "96.4%", "94.8%"],
        ],
        [0.5, 0.25, 0.25],
        facts={
            1: F(
                "zph.net.branches",
                ["Branches", "214", "198"],
                "Zephyra had 214 branches in FY24 (FY23: 198)",
                "table_cell",
            ),
            4: F(
                "zph.net.trucks",
                ["Owned trucks", "2,840", "2,510"],
                "Zephyra had 2,840 owned trucks in FY24 (FY23: 2,510)",
                "table_cell",
            ),
            5: F(
                "zph.net.otd",
                ["On-time delivery rate", "96.4%", "94.8%"],
                "Zephyra's on-time delivery rate was 96.4% in FY24 (FY23: 94.8%)",
                "table_cell",
            ),
        },
    )
    s.text("Branches are spread across 28 states and union territories.", size=15)


def s13(d: Deck) -> None:
    s = d.slide("ESG and people")
    s.bullets(
        [
            (
                "Emissions intensity per tonne-km fell 9 per cent; 1,150 electric vehicles now run in last-mile delivery.",
                (
                    F(
                        "zph.esg.intensity",
                        ["Emissions intensity per tonne-km fell 9 per cent"],
                        "Zephyra's emissions intensity per tonne-km fell 9%",
                        "number",
                    ),
                    F(
                        "zph.esg.ev",
                        ["1,150 electric vehicles"],
                        "Zephyra runs 1,150 electric vehicles in last-mile delivery",
                        "number",
                    ),
                ),
            ),
            (
                "Women made up 14.2 per cent of the workforce of 14,120 employees.",
                (
                    F(
                        "zph.esg.women",
                        ["14.2 per cent of the workforce"],
                        "Women are 14.2% of Zephyra's workforce",
                        "number",
                    ),
                ),
            ),
            (
                "Lost-time injury frequency rate was 0.38 per million man-hours.",
                (
                    F(
                        "zph.esg.ltifr",
                        ["Lost-time injury frequency rate was 0.38"],
                        "Zephyra's LTIFR was 0.38 per million man-hours",
                        "number",
                    ),
                ),
            ),
            (
                "Corporate social responsibility spending was ₹ 11.2 crore.",
                (
                    F(
                        "zph.esg.csr",
                        ["social responsibility spending was ₹ 11.2 crore"],
                        "Zephyra's CSR spend was ₹ 11.2 crore",
                        "number",
                    ),
                ),
            ),
        ],
        size=17,
    )


def s14(d: Deck) -> None:
    s = d.slide("Outlook and FY25 guidance")
    s.bullets(
        [
            (
                "FY25 revenue growth guidance: 12 to 14 per cent.",
                (
                    F(
                        "zph.guide.revenue",
                        ["FY25 revenue growth guidance", "12 to 14 per cent"],
                        "Zephyra guides FY25 revenue growth of 12 to 14%",
                        "number",
                    ),
                ),
            ),
            (
                "FY25 capital expenditure guidance: ₹ 500 to 550 crore, mainly warehouses and electric vehicles.",
                (
                    F(
                        "zph.guide.capex",
                        ["FY25 capital expenditure guidance", "₹ 500 to 550 crore"],
                        "Zephyra guides FY25 capex of ₹ 500 to 550 crore",
                        "number",
                    ),
                ),
            ),
            (
                "Target: net debt to EBITDA below 1.2x by 31 March 2025.",
                (
                    F(
                        "zph.guide.leverage",
                        ["net debt to EBITDA below 1.2x", "31 March 2025"],
                        "Zephyra targets net debt to EBITDA below 1.2x by 31 March 2025",
                        "number",
                    ),
                ),
            ),
            "Priorities: grow contract logistics, scale digital services and keep leverage low.",
        ],
        size=18,
    )


def s15(d: Deck) -> None:
    s = d.slide("Stock information and investor contacts")
    s.table(
        [
            ["Item", "Details"],
            ["ISIN", Z.isin],
            ["BSE scrip code", Z.bse_code],
            ["NSE symbol", Z.nse_symbol],
            ["CIN", Z.cin],
            ["Promoters", "52.4%"],
            ["Foreign portfolio investors", "18.1%"],
            ["Domestic institutions", "15.7%"],
            ["Public", "13.8%"],
            ["Investor relations", "ir@zephyra.example; +91 22 5550 0147"],
        ],
        [0.4, 0.6],
        facts={
            1: F("zph.isin", ["ISIN", Z.isin], f"Zephyra's ISIN is {Z.isin}", "identifier"),
            2: F("zph.bse", ["BSE scrip code", Z.bse_code], f"Zephyra's BSE scrip code is {Z.bse_code}", "identifier"),
            3: F("zph.nse", ["NSE symbol", Z.nse_symbol], f"Zephyra's NSE symbol is {Z.nse_symbol}", "identifier"),
            4: F("zph.cin", ["CIN", Z.cin], f"Zephyra's CIN is {Z.cin}", "identifier"),
            5: F("zph.holding.promoters", ["Promoters", "52.4%"], "Promoters hold 52.4% of Zephyra", "table_cell"),
            9: F(
                "zph.ir",
                ["Investor relations", "ir@zephyra.example", "+91 22 5550 0147"],
                "Zephyra's investor relations contact is ir@zephyra.example, +91 22 5550 0147",
                "table_cell",
            ),
        },
        size=12,
        row_h=0.4,
    )
