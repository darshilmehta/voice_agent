"""DOCX generation (python-docx) that plants facts under the heading they sit in.

Docling reads no page numbers from DOCX, so a fact's locator is its nearest heading (headings are unique within a
document: the builder refuses duplicates). Explicit page breaks give the document its nominal pages for human readers.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt

from .registry import FactSpec, Registry


class DocxDoc:
    def __init__(self, reg: Registry, name: str, *, title: str, font: str | None = None) -> None:
        self.reg = reg
        self.name = name
        self.doc = Document()
        self.doc.core_properties.title = title
        self.doc.core_properties.author = "Eval corpus generator (fictional organisation)"
        self.doc.core_properties.created = datetime(2024, 4, 1)
        self.doc.core_properties.modified = datetime(2024, 4, 1)
        if font:
            style = self.doc.styles["Normal"]
            style.font.name = font
            style.font.size = Pt(11)
        self.current = ""
        self._headings: set[str] = set()
        self.nominal_pages = 1

    def _plant(self, text: str, facts: tuple[FactSpec, ...]) -> None:
        for spec in facts:
            self.reg.plant(spec, document=self.name, text=text, locator=self.current or None)

    def _heading(self, text: str, level: int) -> None:
        text = text.strip()
        if text in self._headings:
            raise ValueError(f"{self.name}: duplicate heading {text!r} (locators must be unique)")
        self._headings.add(text)
        self.current = text
        if level == 0:
            self.doc.add_paragraph(text, style="Title")
        else:
            self.doc.add_heading(text, level=level)

    def title(self, text: str) -> None:
        self._heading(text, 0)

    def h1(self, text: str) -> None:
        self._heading(text, 1)

    def h2(self, text: str) -> None:
        self._heading(text, 2)

    def para(self, text: str, *facts: FactSpec) -> None:
        self.doc.add_paragraph(text)
        self._plant(text, facts)

    def bullets(self, items: list[str | tuple[str, tuple[FactSpec, ...]]]) -> None:
        for item in items:
            text, facts = (item, ()) if isinstance(item, str) else item
            self.doc.add_paragraph(text, style="List Bullet")
            self._plant(text, facts)

    def table(self, rows: list[list[str]], *, facts: dict[int, tuple[FactSpec, ...] | FactSpec] | None = None) -> None:
        table = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r, row in enumerate(rows):
            for c, cell in enumerate(row):
                table.cell(r, c).text = cell
                if r == 0:
                    for run in table.cell(r, c).paragraphs[0].runs:
                        run.bold = True
        for r, specs in (facts or {}).items():
            specs = (specs,) if isinstance(specs, FactSpec) else specs
            row_text = " ".join(rows[r])
            for spec in specs:
                self.reg.plant(spec, document=self.name, text=row_text, locator=self.current or None)
        self.doc.add_paragraph("")

    def page_break(self) -> None:
        self.doc.add_page_break()
        self.nominal_pages += 1

    def save(self, path: Path) -> None:
        self.doc.save(str(path))
