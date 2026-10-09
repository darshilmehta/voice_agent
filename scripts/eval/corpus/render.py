"""Read finished files back into per-page text, to verify the planted facts (pypdf, python-pptx, python-docx)."""

from __future__ import annotations

from pathlib import Path

from app.evals.manifest import RenderedDocument


def render_pdf(path: Path) -> RenderedDocument:
    """Text of every page of a PDF (1-based page n is element n-1)."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return RenderedDocument(pages=[(page.extract_text() or "") for page in reader.pages])


def _shape_texts(shapes) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    out: list[str] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            out += _shape_texts(shape.shapes)
            continue
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            out += [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                out.append(" ".join(cell.text for cell in row.cells))
    return out


def render_pptx(path: Path) -> RenderedDocument:
    """Text of every slide (title, text boxes, tables), slide n at element n-1."""
    from pptx import Presentation

    prs = Presentation(str(path))
    return RenderedDocument(pages=["\n".join(_shape_texts(slide.shapes)) for slide in prs.slides])


def render_docx(path: Path) -> RenderedDocument:
    """Text under each heading of a DOCX (headings are ``Title`` / ``Heading n`` paragraphs; tables count as text of
    the section they sit in). The key is the heading text; text before the first heading is under ``""``."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    sections: dict[str, list[str]] = {"": []}
    current = ""
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            style = para.style.name if para.style is not None else ""
            if style == "Title" or style.startswith("Heading"):
                current = para.text.strip()
                sections.setdefault(current, [])
            elif para.text.strip():
                sections[current].append(para.text)
        elif tag == "tbl":
            table = Table(child, doc)
            for row in table.rows:
                sections[current].append(" ".join(cell.text for cell in row.cells))
    return RenderedDocument(sections={k: "\n".join(v) for k, v in sections.items()})
