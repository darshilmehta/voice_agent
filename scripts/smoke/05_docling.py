# /// script
# requires-python = ">=3.12"
# dependencies = ["docling>=2.50", "rapidocr>=3.9.1,<3.10", "onnxruntime", "reportlab", "python-docx", "python-pptx", "pypdfium2", "pillow", "transformers>=4.45", "opencv-python-headless"]
#
# [tool.uv]
# # docling pulls opencv-python-headless, rapidocr pulls opencv-python: both write the same cv2/ folder,
# # and removing either deletes it. Keep only the headless build (a server never needs OpenCV's GUI).
# override-dependencies = ["opencv-python; sys_platform == 'never'"]
# ///
"""Smoke test 05 — Docling ingestion, fully offline.

Generates test documents into data/smoke/docs/:
  annual_report.pdf   3 pages: headings, prose, a FY23/FY24 financial table (page 2)
  scanned_page.pdf    page 2 rasterized to an image-only PDF (forces OCR)
  hindi_notes.docx    Hindi headings/prose + a table
  board_deck.pptx     2 slides

Then checks: offline conversion with local artifacts, headings, table cells,
page provenance, Devanagari preserved, OCR on the scanned page, seconds/page,
and structure-aware chunking with the BGE-M3 tokenizer (tables not split,
every chunk carries headings + page numbers).

Run:  uv run scripts/smoke/05_docling.py
"""

import os
import resource
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
os.environ["HF_HOME"] = str(ROOT / "data/models/huggingface")
os.environ["HF_HUB_OFFLINE"] = "1"
os.environ["TRANSFORMERS_OFFLINE"] = "1"
os.environ["HF_HUB_DISABLE_TELEMETRY"] = "1"
os.environ.setdefault("TOKENIZERS_PARALLELISM", "false")

DOCS = ROOT / "data/smoke/docs"
ARTIFACTS = ROOT / "data/models/docling"
results: list[tuple[str, bool, str]] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    results.append((name, ok, detail))
    print(f"  {'PASS' if ok else 'FAIL'}  {name}" + (f"  — {detail}" if detail else ""))


def local_snapshot(repo_id: str) -> str:
    snaps = Path(os.environ["HF_HOME"]) / "hub" / f"models--{repo_id.replace('/', '--')}" / "snapshots"
    return str(sorted(snaps.iterdir(), key=lambda p: p.stat().st_mtime)[-1])


# ---------------------------------------------------------------- fixtures

PROSE = ("The company delivered a strong year with growth across all major segments. Management continued to invest "
         "in product development while keeping operating costs under control. ") * 4


def make_pdf(path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    st = getSampleStyleSheet()
    table = Table([
        ["Metric", "FY23", "FY24"],
        ["Revenue from operations (₹ crore)".replace("₹", "INR"), "3,142", "4,210"],
        ["EBITDA (INR crore)", "531", "766"],
        ["EBITDA margin", "16.9%", "18.2%"],
        ["Net profit (INR crore)", "298", "441"],
    ])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                               ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)]))
    story = [
        Paragraph("Annual Report 2024", st["Title"]),
        Paragraph("Chairman's Message", st["Heading1"]), Paragraph(PROSE, st["BodyText"]),
        Paragraph("Company Overview", st["Heading1"]), Paragraph(PROSE, st["BodyText"]),
        PageBreak(),
        Paragraph("Financial Results", st["Heading1"]),
        Paragraph("Key Metrics", st["Heading2"]),
        Paragraph("The table below summarises the key financial metrics for the last two years.", st["BodyText"]),
        Spacer(1, 12), table, Spacer(1, 12),
        Paragraph("Revenue from operations grew 34% year on year, led by the enterprise segment. "
                  "EBITDA margin improved to 18.2% from 16.9% on operating leverage.", st["BodyText"]),
        PageBreak(),
        Paragraph("Risk Management", st["Heading1"]),
        Paragraph("Product SKU-48213 depends on a single supplier. " + PROSE, st["BodyText"]),
    ]
    SimpleDocTemplate(str(path), pagesize=A4).build(story)


def make_scanned(src: Path, path: Path) -> None:
    import pypdfium2 as pdfium
    page = pdfium.PdfDocument(str(src))[1]  # page 2: the financial table
    img = page.render(scale=2.5).to_pil().convert("RGB")
    img.save(path, "PDF", resolution=180)


def make_docx(path: Path) -> None:
    from docx import Document
    d = Document()
    d.add_heading("वार्षिक रिपोर्ट सारांश", level=1)
    d.add_paragraph("कंपनी ने वित्त वर्ष 2024 में अपने कार्बन उत्सर्जन में 15 प्रतिशत की कमी की, जिसका मुख्य कारण सौर ऊर्जा का उपयोग था।")
    d.add_heading("वित्तीय परिणाम", level=2)
    t = d.add_table(rows=3, cols=3)
    for r, row in enumerate([["मापदंड", "FY23", "FY24"], ["राजस्व (करोड़)", "3,142", "4,210"], ["EBITDA मार्जिन", "16.9%", "18.2%"]]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    d.save(path)


def make_pptx(path: Path) -> None:
    from pptx import Presentation
    p = Presentation()
    s1 = p.slides.add_slide(p.slide_layouts[0])
    s1.shapes.title.text = "FY24 Board Update"
    s1.placeholders[1].text = "Investor presentation, April 2025"
    s2 = p.slides.add_slide(p.slide_layouts[1])
    s2.shapes.title.text = "Highlights"
    s2.placeholders[1].text = "Revenue up 34% to INR 4,210 crore\nEBITDA margin 18.2%\nCarbon emissions down 15%"
    p.save(path)


# ---------------------------------------------------------------- test

def main() -> int:
    DOCS.mkdir(parents=True, exist_ok=True)
    pdf, scanned, docx, pptx = DOCS / "annual_report.pdf", DOCS / "scanned_page.pdf", DOCS / "hindi_notes.docx", DOCS / "board_deck.pptx"
    make_pdf(pdf); make_scanned(pdf, scanned); make_docx(docx); make_pptx(pptx)
    print(f"  fixtures written to {DOCS}")

    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions, RapidOcrOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption
    from docling_core.types.doc import DocItemLabel, TableItem

    opts = PdfPipelineOptions(artifacts_path=str(ARTIFACTS))
    opts.do_ocr = True
    opts.do_table_structure = True
    # Explicit engine: auto-detection silently runs with no OCR when it finds none.
    # RapidOCR models are in artifacts_path; rapidocr 3.10 dropped an API docling needs (pin <3.10).
    opts.ocr_options = RapidOcrOptions()
    conv = DocumentConverter(format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=opts)})
    print(f"  OCR engine: {type(opts.ocr_options).__name__}")

    print("\n[1] PDF: annual_report.pdf (3 pages)")
    t0 = time.perf_counter()
    doc = conv.convert(pdf).document
    dt = time.perf_counter() - t0
    pages = len(doc.pages)
    check("converted offline with local artifacts", pages == 3, f"{pages} pages in {dt:.1f}s ({dt / max(pages, 1):.1f} s/page, includes model load)")
    t0 = time.perf_counter(); conv.convert(pdf); warm = time.perf_counter() - t0
    check("warm conversion speed", True, f"{warm / pages:.2f} s/page")
    headings = [it.text for it, _ in doc.iterate_items() if getattr(it, "label", None) in (DocItemLabel.SECTION_HEADER, DocItemLabel.TITLE)]
    check("headings detected", {"Financial Results", "Key Metrics", "Risk Management"} <= set(headings), f"{headings}")
    tables = [t for t in doc.tables if isinstance(t, TableItem)]
    check("one table found", len(tables) == 1, f"{len(tables)} tables")
    if tables:
        df = tables[0].export_to_dataframe(doc=doc)
        flat = df.astype(str).to_string()
        row = df[df.iloc[:, 0].astype(str).str.contains("EBITDA margin")]
        check("table cells correct (EBITDA margin FY24 = 18.2%)", not row.empty and "18.2%" in row.astype(str).to_string(), flat.replace("\n", " | ")[:160])
        check("table on page 2", tables[0].prov[0].page_no == 2, f"page {tables[0].prov[0].page_no}")
    md = doc.export_to_markdown()
    check("markdown keeps the table", "| EBITDA margin" in md and "18.2%" in md)

    print("\n[2] Scanned PDF (image only → OCR)")
    t0 = time.perf_counter()
    sdoc = conv.convert(scanned).document
    smd = sdoc.export_to_markdown()
    check("OCR recovered text", "Financial Results" in smd and "18.2" in smd, f"{time.perf_counter() - t0:.1f}s, {len(smd)} chars")
    check("OCR recovered the table", len(sdoc.tables) >= 1, f"{len(sdoc.tables)} tables")

    print("\n[3] DOCX (Hindi)")
    ddoc = conv.convert(docx).document
    dmd = ddoc.export_to_markdown()
    check("Devanagari preserved", "कार्बन उत्सर्जन" in dmd and "वित्तीय परिणाम" in dmd)
    check("Hindi table parsed", len(ddoc.tables) == 1 and "18.2%" in dmd)

    print("\n[4] PPTX")
    pmd = conv.convert(pptx).document.export_to_markdown()
    check("slide text extracted", "FY24 Board Update" in pmd and "18.2%" in pmd)

    print("\n[5] Chunking (HybridChunker + BGE-M3 tokenizer, max 512 tokens)")
    from docling.chunking import HybridChunker
    from docling_core.transforms.chunker.tokenizer.huggingface import HuggingFaceTokenizer
    from transformers import AutoTokenizer
    tok = HuggingFaceTokenizer(tokenizer=AutoTokenizer.from_pretrained(local_snapshot("BAAI/bge-m3")), max_tokens=512)
    chunker = HybridChunker(tokenizer=tok, merge_peers=True)
    chunks = list(chunker.chunk(dl_doc=doc))
    for ch in chunks:
        pg = sorted({p.page_no for it in ch.meta.doc_items for p in it.prov})
        print(f"        pages {pg} headings {ch.meta.headings} tokens {tok.count_tokens(chunker.contextualize(ch))}  {ch.text[:50]!r}")
    check("every chunk has page numbers", all(any(it.prov for it in ch.meta.doc_items) for ch in chunks))
    check("every chunk has a heading path", all(ch.meta.headings for ch in chunks))
    check("chunks ≤ 512 tokens", all(tok.count_tokens(chunker.contextualize(ch)) <= 512 for ch in chunks))
    table_chunks = [ch for ch in chunks if any(it.label == DocItemLabel.TABLE for it in ch.meta.doc_items)]
    check("table kept whole in one chunk", len(table_chunks) == 1 and "18.2%" in table_chunks[0].text and "4,210" in table_chunks[0].text)

    print("\n[6] Memory")
    check("peak RSS", True, f"{resource.getrusage(resource.RUSAGE_SELF).ru_maxrss / 1e9:.2f} GB")

    failed = [r for r in results if not r[1]]
    print(f"\n{len(results) - len(failed)}/{len(results)} checks passed")
    return 1 if failed else 0


if __name__ == "__main__":
    code = main()
    sys.stdout.flush(); sys.stderr.flush()
    os._exit(code)  # skip interpreter teardown: onnxruntime threads abort on macOS during shutdown (libc++abi recursive_mutex)
